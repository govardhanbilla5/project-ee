"""
Batch forecast service.

Reads (Part Number, Tier 1) pairs from an uploaded Excel, runs the
forecast engine for each unique combination, and writes a result
workbook with one sheet per forecast month (12 sheets, or 13 when the
current month is included) plus a Summary.

COLUMN LAYOUT (per monthly sheet)
──────────────────────────────────
Part Number | Tier 1 | Weight (lbs) | Base Price ($) |
Quarter (Current) | Quarter (Previous) |
MC_Q ($/lb) | MC_Q-1 ($/lb) |
PPI_Q | PPI_Q-1 | PPI Factor |
CNG_Q ($/lb) | CNG_Q-1 ($/lb) |
AMS_Q ($/lb) | AMS_Q-1 ($/lb) | AMS Delta ($/lb) |
DF_c | Predicted Price ($)
"""

import io
import logging
from typing import Optional

import openpyxl
import pandas as pd
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app.models.response import ForecastResponse
from app.services.forecast_engine import ForecastEngine

logger = logging.getLogger(__name__)

# ─────────────────────────────────────────────────────────────────────────────
# STYLING
# ─────────────────────────────────────────────────────────────────────────────

_THIN        = Side(style="thin", color="BFBFBF")
_BORDER      = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
_ALT_FILL    = PatternFill("solid", fgColor="D6E4F0")
_WHITE_FILL  = PatternFill("solid", fgColor="FFFFFF")
_WARN_FILL   = PatternFill("solid", fgColor="FCE4D6")
_HEADER_FONT = Font(name="Arial", bold=True, color="FFFFFF", size=10)
_DATA_FONT   = Font(name="Arial", size=10)
_WARN_FONT   = Font(name="Arial", size=10, color="C00000")
_CENTER      = Alignment(horizontal="center", vertical="center", wrap_text=True)
_LEFT        = Alignment(horizontal="left",   vertical="center")


# ─────────────────────────────────────────────────────────────────────────────
# COLUMN DEFINITIONS
# ─────────────────────────────────────────────────────────────────────────────

def _build_columns() -> list[tuple[str, str, callable]]:
    def _ctx(mf): return mf.quarter_context
    return [
        ("Part Number",          "@",             lambda pn, t1, fr, mf: pn),
        ("Tier 1",               "@",             lambda pn, t1, fr, mf: t1),
        ("Weight\n(lbs)",        "0.00",          lambda pn, t1, fr, mf: fr.pwt_lbs),
        ("Base Price\n($)",      "$#,##0.0000",   lambda pn, t1, fr, mf: mf.base_price_used),
        ("Quarter\n(Current)",   "@",             lambda pn, t1, fr, mf: _ctx(mf).quarter_label),
        ("Quarter\n(Previous)",  "@",             lambda pn, t1, fr, mf: _ctx(mf).prev_quarter_label),
        ("MC_Q\n($/lb)",         "$0.000000",     lambda pn, t1, fr, mf: _ctx(mf).mc_q),
        ("MC_Q-1\n($/lb)",       "$0.000000",     lambda pn, t1, fr, mf: _ctx(mf).mc_q_1),
        ("PPI_Q",                "0.000",         lambda pn, t1, fr, mf: _ctx(mf).ppi_q),
        ("PPI_Q-1",              "0.000",         lambda pn, t1, fr, mf: _ctx(mf).ppi_q_1),
        ("PPI Factor",           "0.000000%",     lambda pn, t1, fr, mf: _ctx(mf).ppi_factor),
        ("CNG_Q\n($/lb)",        "$0.0000",       lambda pn, t1, fr, mf: _ctx(mf).cng_q),
        ("CNG_Q-1\n($/lb)",      "$0.0000",       lambda pn, t1, fr, mf: _ctx(mf).cng_q_1),
        ("AMS_Q\n($/lb)",        "$0.000000",     lambda pn, t1, fr, mf: _ctx(mf).ams_q),
        ("AMS_Q-1\n($/lb)",      "$0.000000",     lambda pn, t1, fr, mf: _ctx(mf).ams_q_1),
        ("AMS Delta\n($/lb)",    "$0.000000",     lambda pn, t1, fr, mf: _ctx(mf).ams_delta),
        ("DF_c",                 "0.00",          lambda pn, t1, fr, mf: mf.df_c),
        ("Predicted Price\n($)", "$#,##0.0000",   lambda pn, t1, fr, mf: mf.predicted_price),
    ]


COLUMNS = _build_columns()
N_COLS  = len(COLUMNS)


# ─────────────────────────────────────────────────────────────────────────────
# CELL HELPERS
# ─────────────────────────────────────────────────────────────────────────────

def _write_header(ws, col_idx: int, label: str) -> None:
    cell = ws.cell(row=1, column=col_idx, value=label)
    cell.font = _HEADER_FONT; cell.fill = _HEADER_FILL
    cell.alignment = _CENTER; cell.border = _BORDER


def _write_data(ws, row: int, col: int, value, fmt: str, fill, left=False) -> None:
    cell = ws.cell(row=row, column=col, value=value)
    cell.font = _DATA_FONT; cell.fill = fill; cell.border = _BORDER
    cell.number_format = fmt
    cell.alignment = _LEFT if left else _CENTER


def _write_error_row(ws, row: int, part_number: str, tier_1: str, error_msg: str, fill) -> None:
    for col, val in [(1, part_number), (2, tier_1)]:
        c = ws.cell(row=row, column=col, value=val)
        c.font = _WARN_FONT; c.fill = _WARN_FILL
        c.border = _BORDER; c.alignment = _LEFT

    err = ws.cell(row=row, column=3, value=f"ERROR: {error_msg}")
    err.font = _WARN_FONT; err.fill = _WARN_FILL
    err.border = _BORDER; err.alignment = _LEFT
    ws.merge_cells(start_row=row, start_column=3, end_row=row, end_column=N_COLS)


def _set_col_widths(ws) -> None:
    widths = [22, 22, 10, 14, 14, 14, 13, 13, 10, 10, 13, 13, 13, 13, 13, 13, 8, 18]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[1].height = 36


# ─────────────────────────────────────────────────────────────────────────────
# INPUT READER
# ─────────────────────────────────────────────────────────────────────────────

def read_parts_from_upload(file_bytes: bytes) -> list[tuple[str, str]]:
    """
    Read (Part Number, Tier 1) pairs from the uploaded Excel file.

    Expects columns named 'Part Number' and 'Tier 1' (case-insensitive).
    Returns a deduplicated list of (part_number, tier_1) tuples,
    preserving input order.
    """
    df = pd.read_excel(
        io.BytesIO(file_bytes),
        engine="openpyxl",
        dtype=str,
        keep_default_na=False,
        na_values=[""],
    )

    cols_lower = {c.strip().lower(): c for c in df.columns}

    if "part number" not in cols_lower:
        raise ValueError(
            f"Input Excel must have a 'Part Number' column. "
            f"Found: {list(df.columns)}"
        )
    if "tier 1" not in cols_lower:
        raise ValueError(
            f"Input Excel must have a 'Tier 1' column. "
            f"Found: {list(df.columns)}"
        )

    pn_col = cols_lower["part number"]
    t1_col = cols_lower["tier 1"]

    seen: set[tuple[str, str]] = set()
    pairs: list[tuple[str, str]] = []
    for _, row in df.iterrows():
        pn = str(row[pn_col]).strip()
        t1 = str(row[t1_col]).strip()
        if pn and t1 and (pn, t1) not in seen:
            seen.add((pn, t1))
            pairs.append((pn, t1))

    logger.info("Read %d unique (part, tier_1) pairs from uploaded file", len(pairs))
    return pairs


# CNG column name aliases (case-insensitive, matched against stripped headers).
_CNG_Q_ALIASES  = ("cng_q", "cng q", "cng_q current", "cng current", "cng - current qtr")
_CNG_Q1_ALIASES = ("cng_q-1", "cng_q1", "cng q-1", "cng_q_1", "cng previous", "cng - previous qtr")


def read_parts_with_cng_from_upload(file_bytes: bytes) -> list[dict]:
    """
    Read (Part Number, Tier 1, CNG_Q, CNG_Q-1) rows from the uploaded Excel.

    The CNG values are taken straight from the file — the caller no longer
    supplies them. Expects columns 'Part Number', 'Tier 1', 'CNG_Q' and
    'CNG_Q-1' (case-insensitive; a few common spellings are accepted).
    Returns a deduplicated list of dicts, preserving input order:
        {"part_number", "tier_1", "cng_q", "cng_q_1"}
    """
    df = pd.read_excel(
        io.BytesIO(file_bytes),
        engine="openpyxl",
        dtype=str,
        keep_default_na=False,
        na_values=[""],
    )

    cols_lower = {c.strip().lower(): c for c in df.columns}

    if "part number" not in cols_lower:
        raise ValueError(
            f"Input Excel must have a 'Part Number' column. Found: {list(df.columns)}"
        )
    if "tier 1" not in cols_lower:
        raise ValueError(
            f"Input Excel must have a 'Tier 1' column. Found: {list(df.columns)}"
        )

    def _find(aliases: tuple[str, ...]) -> Optional[str]:
        for a in aliases:
            if a in cols_lower:
                return cols_lower[a]
        return None

    cng_q_col  = _find(_CNG_Q_ALIASES)
    cng_q1_col = _find(_CNG_Q1_ALIASES)
    if cng_q_col is None or cng_q1_col is None:
        raise ValueError(
            "Input Excel must have 'CNG_Q' and 'CNG_Q-1' columns "
            f"(CNG is now read from the file). Found: {list(df.columns)}"
        )

    pn_col = cols_lower["part number"]
    t1_col = cols_lower["tier 1"]

    seen: set[tuple[str, str]] = set()
    rows: list[dict] = []
    for _, row in df.iterrows():
        pn = str(row[pn_col]).strip()
        t1 = str(row[t1_col]).strip()
        if not (pn and t1) or (pn, t1) in seen:
            continue
        try:
            cng_q   = float(row[cng_q_col])
            cng_q_1 = float(row[cng_q1_col])
        except (ValueError, TypeError):
            raise ValueError(
                f"Non-numeric CNG value for part '{pn}' / '{t1}'. "
                "CNG_Q and CNG_Q-1 must be numbers."
            )
        seen.add((pn, t1))
        rows.append({"part_number": pn, "tier_1": t1, "cng_q": cng_q, "cng_q_1": cng_q_1})

    logger.info("Read %d unique (part, tier_1, cng) rows from uploaded file", len(rows))
    return rows


# ─────────────────────────────────────────────────────────────────────────────
# OUTPUT BUILDER
# ─────────────────────────────────────────────────────────────────────────────

def _run_forecasts_uniform(
    part_tier_pairs: list[tuple[str, str]],
    engine: ForecastEngine,
    cng_q: float,
    cng_q_1: float,
    include_current_month: bool,
) -> dict[tuple[str, str], "ForecastResponse | Exception"]:
    """Run the engine for every pair using ONE shared CNG_Q / CNG_Q-1."""
    results: dict[tuple[str, str], ForecastResponse | Exception] = {}
    for pn, t1 in part_tier_pairs:
        try:
            results[(pn, t1)] = engine.forecast(
                part_number=pn, tier_1=t1, cng_q=cng_q, cng_q_1=cng_q_1,
                include_current_month=include_current_month,
            )
            logger.debug("Forecast OK: %s / %s", pn, t1)
        except Exception as exc:
            results[(pn, t1)] = exc
            logger.warning("Forecast failed for %s / %s: %s", pn, t1, exc)
    return results


def _run_forecasts_per_row(
    part_rows: list[dict],
    engine: ForecastEngine,
    include_current_month: bool,
) -> tuple[list[tuple[str, str]], dict[tuple[str, str], "ForecastResponse | Exception"]]:
    """Run the engine per row, each with its OWN CNG_Q / CNG_Q-1 from the file."""
    results: dict[tuple[str, str], ForecastResponse | Exception] = {}
    part_tier_pairs: list[tuple[str, str]] = []
    for row in part_rows:
        key = (row["part_number"], row["tier_1"])
        part_tier_pairs.append(key)
        try:
            results[key] = engine.forecast(
                part_number=key[0], tier_1=key[1],
                cng_q=row["cng_q"], cng_q_1=row["cng_q_1"],
                include_current_month=include_current_month,
            )
            logger.debug("Forecast OK: %s / %s (CNG_Q=%s)", key[0], key[1], row["cng_q"])
        except Exception as exc:
            results[key] = exc
            logger.warning("Forecast failed for %s / %s: %s", key[0], key[1], exc)
    return part_tier_pairs, results


def build_forecast_workbook(
    part_tier_pairs: list[tuple[str, str]],
    engine: ForecastEngine,
    cng_q: float,
    cng_q_1: float,
    include_current_month: bool = False,
) -> bytes:
    """
    Run forecasts for all (part_number, tier_1) pairs with a single shared
    CNG pair and build the output workbook. Kept for backward compatibility.
    """
    results = _run_forecasts_uniform(
        part_tier_pairs, engine, cng_q, cng_q_1, include_current_month
    )
    return _build_workbook_from_results(part_tier_pairs, results)


def build_forecast_workbook_from_rows(
    part_rows: list[dict],
    engine: ForecastEngine,
    include_current_month: bool = False,
) -> bytes:
    """Run forecasts using per-row CNG (from the file) and build the workbook."""
    part_tier_pairs, results = _run_forecasts_per_row(
        part_rows, engine, include_current_month
    )
    return _build_workbook_from_results(part_tier_pairs, results)


def build_preview_from_rows(
    part_rows: list[dict],
    engine: ForecastEngine,
    include_current_month: bool = False,
) -> dict:
    """
    Run forecasts using per-row CNG and return a JSON-serialisable preview
    (Summary + one entry per monthly sheet) for on-screen display.
    """
    part_tier_pairs, results = _run_forecasts_per_row(
        part_rows, engine, include_current_month
    )

    month_labels: list[tuple[str, str]] = []
    for r in results.values():
        if isinstance(r, ForecastResponse):
            month_labels = [(f.year_month, f.month_label) for f in r.forecasts]
            break
    if not month_labels:
        first_error = next(iter(results.values()), None)
        raise ValueError(
            f"All parts failed forecasting — nothing to preview. First error: {first_error}"
        )

    n_ok = sum(1 for r in results.values() if isinstance(r, ForecastResponse))
    n_failed = len(results) - n_ok

    summary: list[dict] = []
    for pn, t1 in part_tier_pairs:
        r = results[(pn, t1)]
        if isinstance(r, Exception):
            summary.append({"part_number": pn, "tier_1": t1, "error": str(r)})
            continue
        monthly = []
        for ym, _ in month_labels:
            mf = next((f for f in r.forecasts if f.year_month == ym), None)
            monthly.append(mf.predicted_price if mf else None)
        ctx0 = r.forecasts[0].quarter_context
        summary.append({
            "part_number": pn, "tier_1": t1,
            "weight_lbs": r.pwt_lbs, "base_price": r.base_price,
            "cng_q": ctx0.cng_q, "cng_q_1": ctx0.cng_q_1,
            "monthly": monthly, "error": None,
        })

    monthly_sheets: list[dict] = []
    for ym, label in month_labels:
        rows: list[dict] = []
        for pn, t1 in part_tier_pairs:
            r = results[(pn, t1)]
            if isinstance(r, Exception):
                rows.append({"part_number": pn, "tier_1": t1, "error": str(r)})
                continue
            mf = next((f for f in r.forecasts if f.year_month == ym), None)
            if mf is None:
                rows.append({"part_number": pn, "tier_1": t1, "error": f"No data for {ym}"})
                continue
            q = mf.quarter_context
            rows.append({
                "part_number": pn, "tier_1": t1,
                "weight_lbs": r.pwt_lbs, "base_price_used": mf.base_price_used,
                "quarter": q.quarter_label,
                "mc_q": q.mc_q, "mc_q_1": q.mc_q_1,
                "ppi_q": q.ppi_q, "ppi_q_1": q.ppi_q_1, "ppi_factor": q.ppi_factor,
                "cng_q": q.cng_q, "cng_q_1": q.cng_q_1,
                "ams_q": q.ams_q, "ams_q_1": q.ams_q_1, "ams_delta": q.ams_delta,
                "predicted_price": mf.predicted_price, "error": None,
            })
        monthly_sheets.append({"year_month": ym, "label": label, "rows": rows})

    return {
        "part_count": len(part_tier_pairs),
        "ok": n_ok,
        "failed": n_failed,
        "months": [{"year_month": ym, "month_label": label} for ym, label in month_labels],
        "summary": summary,
        "monthly_sheets": monthly_sheets,
    }


def _build_workbook_from_results(
    part_tier_pairs: list[tuple[str, str]],
    results: dict[tuple[str, str], "ForecastResponse | Exception"],
) -> bytes:
    """
    Build the formatted .xlsx workbook from already-computed results.

    include_current_month affects the number of monthly sheets via the
    supplied results. Failed rows show an ERROR message instead of stopping
    the batch. Returns raw bytes of the generated .xlsx workbook.
    """
    # ── Step 2: determine month labels from first successful result ───────
    month_labels: list[tuple[str, str]] = []
    for r in results.values():
        if isinstance(r, ForecastResponse):
            month_labels = [(f.year_month, f.month_label) for f in r.forecasts]
            break

    if not month_labels:
        first_error = next(iter(results.values()), None)
        raise ValueError(
            "All parts failed forecasting — cannot generate output workbook. "
            f"First error: {first_error}"
        )

    # ── Step 3: build workbook ────────────────────────────────────────────
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    for year_month, month_label in month_labels:
        ws = wb.create_sheet(title=month_label[:31])
        ws.freeze_panes = "A2"

        for ci, (label, _, _) in enumerate(COLUMNS, 1):
            _write_header(ws, ci, label)

        for ri, (pn, t1) in enumerate(part_tier_pairs, 2):
            fill   = _ALT_FILL if ri % 2 == 0 else _WHITE_FILL
            result = results[(pn, t1)]

            if isinstance(result, Exception):
                _write_error_row(ws, ri, pn, t1, str(result), fill)
                continue

            mf = next((f for f in result.forecasts if f.year_month == year_month), None)
            if mf is None:
                _write_error_row(ws, ri, pn, t1, f"No forecast data for {year_month}", fill)
                continue

            for ci, (_, fmt, extractor) in enumerate(COLUMNS, 1):
                left = ci in (1, 2)
                try:
                    value = extractor(pn, t1, result, mf)
                except Exception as ex:
                    value = f"ERR: {ex}"
                _write_data(ws, ri, ci, value, fmt, fill, left=left)

        _set_col_widths(ws)

    # ── Step 4: summary sheet ─────────────────────────────────────────────
    _build_summary_sheet(wb, part_tier_pairs, results, month_labels)

    # ── Step 5: serialise ─────────────────────────────────────────────────
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read()


def _build_summary_sheet(
    wb, part_tier_pairs, results, month_labels
) -> None:
    ws = wb.create_sheet(title="Summary", index=0)
    ws.freeze_panes = "E2"

    fixed   = ["Part Number", "Tier 1", "Weight (lbs)", "Base Price ($)"]
    monthly = [label for _, label in month_labels]

    for ci, h in enumerate(fixed + monthly, 1):
        _write_header(ws, ci, h)
    ws.row_dimensions[1].height = 36

    for ri, (pn, t1) in enumerate(part_tier_pairs, 2):
        fill   = _ALT_FILL if ri % 2 == 0 else _WHITE_FILL
        result = results[(pn, t1)]

        if isinstance(result, Exception):
            _write_error_row(ws, ri, pn, t1, str(result), fill)
            continue

        _write_data(ws, ri, 1, pn,               "@",            fill, left=True)
        _write_data(ws, ri, 2, t1,               "@",            fill, left=True)
        _write_data(ws, ri, 3, result.pwt_lbs,   "0.00",         fill)
        _write_data(ws, ri, 4, result.base_price, "$#,##0.0000", fill)

        for mi, (year_month, _) in enumerate(month_labels):
            mf  = next((f for f in result.forecasts if f.year_month == year_month), None)
            val = mf.predicted_price if mf else "N/A"
            _write_data(ws, ri, 5 + mi, val, "$#,##0.0000", fill)

    # Column widths
    for col, width in zip("ABCD", [22, 22, 12, 16]):
        ws.column_dimensions[col].width = width
    for i in range(len(month_labels)):
        ws.column_dimensions[get_column_letter(5 + i)].width = 18
